#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl"]
# ///
"""
Convert XLSX objective templates into hierarchical JSON.

Parses the Phase → Primary → Secondary → Tertiary structure
into a nested JSON tree suitable for engineering documentation.
"""

import json
import sys
from pathlib import Path

import openpyxl


def parse_sheet_to_hierarchy(ws):
    """Parse a worksheet into a hierarchical structure.

    Returns a list of phases, each containing nested objectives.
    """
    phases = []
    current_phase = None
    current_primary = None
    current_secondary = None

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=4):
        a = str(row[0].value).strip() if row[0].value else ""
        b = str(row[1].value).strip() if row[1].value else ""
        c = str(row[2].value).strip() if len(row) > 2 and row[2].value else ""
        d = str(row[3].value).strip() if len(row) > 3 and row[3].value else ""

        # Skip header rows and empty rows
        if a in ("PHASE", "") and b in ("PRIMARY OBJECTIVE", "") and not c and not d:
            if not a and not b:
                continue
            if a == "PHASE":
                continue

        # Skip title rows (first row with template name)
        if a and not b and not c and not d:
            # Check if it's a phase name (known phases or new phase)
            if current_phase is not None or a in get_known_phases():
                current_phase = {
                    "phase": a,
                    "objectives": []
                }
                phases.append(current_phase)
                current_primary = None
                current_secondary = None
            # else: it's a title row, skip
            continue

        # Primary objective (column B)
        if b and not c and not d:
            if current_phase is None:
                continue
            current_primary = {
                "name": b,
                "subObjectives": []
            }
            current_phase["objectives"].append(current_primary)
            current_secondary = None
            continue

        # Secondary objective (column C)
        if c and not d:
            if current_primary is None:
                continue
            current_secondary = {
                "name": c,
                "subObjectives": []
            }
            current_primary["subObjectives"].append(current_secondary)
            continue

        # Tertiary objective (column D)
        if d:
            if current_secondary is None:
                continue
            current_secondary["subObjectives"].append({
                "name": d
            })
            continue

    return phases


def get_known_phases():
    """Return the set of known phase names."""
    return {
        "Project Initiation",
        "Acquisitions & Investment",
        "Licensing, Registrations & Certifications",
        "Insurance",
        "Legal",
        "Finance",
        "Environmental",
        "Community Engagement & Communication",
        "Design & Development",
        "Regulatory Approvals & Permits",
        "Bidding",
        "Construction",
        "Product Procurement & Supply Chain",
        "Workforce Management & Tracking",
        "Commissioning & Handover",
        "Post-Construction/Closeout",
        "Maintenance",
        "Decommissioning",
        "Sale of Property",
    }


def clean_hierarchy(phases):
    """Remove empty subObjectives arrays for cleaner output."""
    cleaned = []
    for phase in phases:
        clean_phase = {
            "phase": phase["phase"],
            "objectives": clean_objectives(phase["objectives"])
        }
        if clean_phase["objectives"]:
            cleaned.append(clean_phase)
    return cleaned


def clean_objectives(objectives):
    """Recursively clean objectives."""
    result = []
    for obj in objectives:
        clean_obj = {"name": obj["name"]}
        if "subObjectives" in obj and obj["subObjectives"]:
            clean_obj["subObjectives"] = clean_objectives(obj["subObjectives"])
        result.append(clean_obj)
    return result


def count_objectives(phases):
    """Count total objectives at all levels."""
    total = 0
    for phase in phases:
        total += count_obj_list(phase.get("objectives", []))
    return total


def count_obj_list(objectives):
    count = len(objectives)
    for obj in objectives:
        count += count_obj_list(obj.get("subObjectives", []))
    return count


def main():
    if len(sys.argv) < 2:
        print("Usage: xlsx_to_hierarchy.py <file.xlsx> [--sheet <name>] [--all] [--stats]", file=sys.stderr)
        sys.exit(1)

    filepath = Path(sys.argv[1])
    if not filepath.exists():
        print(f"Error: File '{filepath}' not found.", file=sys.stderr)
        sys.exit(1)

    sheet_name = None
    show_all = False
    stats_only = False

    i = 2
    while i < len(sys.argv):
        if sys.argv[i] == "--sheet" and i + 1 < len(sys.argv):
            sheet_name = sys.argv[i + 1]
            i += 2
        elif sys.argv[i] == "--all":
            show_all = True
            i += 1
        elif sys.argv[i] == "--stats":
            stats_only = True
            i += 1
        else:
            i += 1

    wb = openpyxl.load_workbook(str(filepath), data_only=True, read_only=True)

    # Skip "Project Types" sheet (index sheet)
    template_sheets = [s for s in wb.sheetnames if s != "Project Types"]

    if sheet_name:
        if sheet_name not in wb.sheetnames:
            print(f"Error: Sheet '{sheet_name}' not found.", file=sys.stderr)
            sys.exit(2)
        template_sheets = [sheet_name]
    elif not show_all:
        template_sheets = [template_sheets[0]]

    results = {}
    for name in template_sheets:
        ws = wb[name]
        hierarchy = parse_sheet_to_hierarchy(ws)
        cleaned = clean_hierarchy(hierarchy)
        results[name] = cleaned

    if stats_only:
        for name, hierarchy in results.items():
            phases_count = len(hierarchy)
            obj_count = count_objectives(hierarchy)
            print(f"{name}: {phases_count} phases, {obj_count} objectives")
        wb.close()
        return

    if len(results) == 1:
        name = list(results.keys())[0]
        print(json.dumps(results[name], indent=2, ensure_ascii=False))
    else:
        print(json.dumps(results, indent=2, ensure_ascii=False))

    wb.close()


if __name__ == "__main__":
    main()
